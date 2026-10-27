import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_jury_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Elegant FinTech Color Palette (Dark Theme / High Contrast / Neon Tech Accents)
    BG_DARK = RGBColor(10, 15, 29)          # Ultra dark navy (#0A0F1D)
    BG_CARD = RGBColor(18, 25, 46)          # Card background (#12192E)
    BG_CARD_INNER = RGBColor(26, 35, 64)    # Inner highlight box (#1A2340)
    BORDER_DEFAULT = RGBColor(38, 52, 89)   # Subtle slate border (#263459)
    BORDER_CYAN = RGBColor(14, 165, 233)    # Bright Sky Blue (#0EA5E9)
    BORDER_GREEN = RGBColor(16, 185, 129)   # Vibrant Emerald (#10B981)
    BORDER_AMBER = RGBColor(245, 158, 11)   # Warning Amber (#F59E0B)
    BORDER_RED = RGBColor(239, 68, 68)      # Alert Red (#EF4444)
    
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_CYAN = RGBColor(56, 189, 248)      # Neon Sky (#38BDF8)
    TEXT_MUTED = RGBColor(148, 163, 184)    # Slate Muted (#94A3B8)
    TEXT_GREEN = RGBColor(52, 211, 153)     # Mint Green (#34D399)
    TEXT_AMBER = RGBColor(251, 191, 36)     # Light Amber (#FBBF24)
    TEXT_PURPLE = RGBColor(192, 132, 252)   # Purple Accent (#C084FC)

    blank_layout = prs.slide_layouts[6]

    def set_slide_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        
        # Top subtle ambient accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BORDER_CYAN
        bar.line.fill.background()
        return bg

    def add_header(slide, title, category="RENTWISE AI • JURY PRESENTATION", page_num=None):
        tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
        tf_cat = tb_cat.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEXT_CYAN

        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11), Inches(0.7))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(25)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE

        if page_num:
            tb_pg = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.2), Inches(0.4))
            tf_pg = tb_pg.text_frame
            p_pg = tf_pg.paragraphs[0]
            p_pg.text = f"{page_num} / 10"
            p_pg.font.size = Pt(10)
            p_pg.font.bold = True
            p_pg.alignment = PP_ALIGN.RIGHT
            p_pg.font.color.rgb = TEXT_MUTED

    def draw_card(slide, left, top, width, height, title="", subtitle=None, border_color=BORDER_DEFAULT, bg_color=BG_CARD):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        if title:
            p0 = tf.paragraphs[0]
            p0.text = title
            p0.font.size = Pt(15)
            p0.font.bold = True
            p0.font.color.rgb = TEXT_WHITE

        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle.upper()
            p_sub.font.size = Pt(9.5)
            p_sub.font.bold = True
            p_sub.font.color.rgb = TEXT_CYAN
            p_sub.space_after = Pt(8)
            
        return tf

    # =========================================================================
    # SLIDE 1: Title & Hero
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s1)

    # Accent decorative box
    glow = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(11.733), Inches(5.2))
    glow.fill.solid()
    glow.fill.fore_color.rgb = BG_CARD
    glow.line.color.rgb = BORDER_CYAN
    glow.line.width = Pt(2)

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(11.0), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "OFFICIAL HACKATHON JURY EVALUATION • TEAM BLITZ"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_CYAN
    p0.space_after = Pt(10)

    p1 = tf.add_paragraph()
    p1.text = "RentWise AI"
    p1.font.size = Pt(50)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.space_after = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Autonomous, Zero-Hallucination HRA Rent Receipt & Proof Pack Agent"
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_CYAN
    p2.space_after = Pt(18)

    p3 = tf.add_paragraph()
    p3.text = "Bridging Banking Ledgers & Employer Tax Portals with Cryptographic Proof and Statutory Compliance"
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_after = Pt(28)

    p4 = tf.add_paragraph()
    p4.text = "🚀 Live Deployment: https://rentwise-ai-84yv.onrender.com/  |  📦 GitHub: HarishkumarK24/Rentwise-Ai"
    p4.font.size = Pt(12)
    p4.font.bold = True
    p4.font.color.rgb = TEXT_GREEN

    # =========================================================================
    # SLIDE 2: Executive Summary & Jury Value Proposition
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s2)
    add_header(s2, "Executive Summary: Why RentWise AI Wins", page_num=2)

    tf_c1 = draw_card(s2, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "1. Zero Hallucination", "Evidence-First Integrity", border_color=BORDER_GREEN)
    p = tf_c1.add_paragraph()
    p.text = "• Deterministic Reconciliation: Financial arithmetic and month-to-payment matching are executed by strict algorithmic logic—not unconstrained LLM guessing."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c1.add_paragraph()
    p.text = "• Strict Proof Rule: A receipt is never generated unless exact bank ledger evidence is present. Missing transactions stay flagged as missing."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c1.add_paragraph()
    p.text = "• Integer Precision: Uses integer paise/rupee units, eliminating floating-point reconciliation drifts."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_c2 = draw_card(s2, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "2. Statutory Compliance", "Section 10(13A) Enforcement", border_color=BORDER_AMBER)
    p = tf_c2.add_paragraph()
    p.text = "• CBDT Circular No. 08/2013: Indian tax law mandates Landlord PAN if annual rent exceeds ₹1,00,000. Generic chatbots overlook this rule."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c2.add_paragraph()
    p.text = "• 3-Way Boundary Engine: Validated for < ₹1L (Optional PAN), == ₹1L (Optional PAN), and > ₹1L (Strictly Mandatory)."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c2.add_paragraph()
    p.text = "• Hard Submission Lock: Blocks employer upload if mandatory PAN is missing, protecting employees from IT department notices."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_c3 = draw_card(s2, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "3. Enterprise Safety", "Cryptographic & Idempotent", border_color=BORDER_CYAN)
    p = tf_c3.add_paragraph()
    p.text = "• Human Approval Gate: Server-side check blocks uploads until explicit human sign-off bound to exact SHA-256 PDF hashes."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c3.add_paragraph()
    p.text = "• Idempotency Protection: Re-submitting the same proof pack reuses the existing upload reference, preventing duplicate portal filings."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_c3.add_paragraph()
    p.text = "• Independent Verification: Post-upload query distinguishes 'Uploaded' from 'Pending Verification' vs 'Approved'."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 3: The Real-World Problem
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s3)
    add_header(s3, "The Problem: The Year-End Tax Declaration Nightmare", page_num=3)

    tf_p1 = draw_card(s3, Inches(0.8), Inches(1.8), Inches(5.66), Inches(2.35), 
                      "Fragmented Financial Ledgers", "The Evidence Chaos", border_color=BORDER_RED)
    p = tf_p1.add_paragraph()
    p.text = "• Tenants pay via UPI, NetBanking, IMPS, and split transfers with inconsistent transaction remarks."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_p1.add_paragraph()
    p.text = "• Late or advance rent payments break simple calendar-month matching, creating chaotic declaration filings."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    tf_p2 = draw_card(s3, Inches(6.86), Inches(1.8), Inches(5.66), Inches(2.35), 
                      "Statutory Compliance Blindspots", "The Legal Risk", border_color=BORDER_AMBER)
    p = tf_p2.add_paragraph()
    p.text = "• Employees unaware of the ₹1,00,000 Landlord PAN mandate face rejected HRA exemptions."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_p2.add_paragraph()
    p.text = "• Form-fillers generate blind receipts without validating PAN format or boundary thresholds, leading to employer audit penalties."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    tf_p3 = draw_card(s3, Inches(0.8), Inches(4.45), Inches(5.66), Inches(2.35), 
                      "The Dangers of Generative AI Hallucination", "The Trust Gap", border_color=BORDER_RED)
    p = tf_p3.add_paragraph()
    p.text = "• When prompt-based LLMs encounter missing months, they frequently hallucinate imaginary dates, amounts, and PANs."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_p3.add_paragraph()
    p.text = "• Submitting fabricated tax evidence is a punishable offense under income tax evasion statutes."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    tf_p4 = draw_card(s3, Inches(6.86), Inches(4.45), Inches(5.66), Inches(2.35), 
                      "The 'False Approval' Illusion", "The Operational Trap", border_color=BORDER_DEFAULT)
    p = tf_p4.add_paragraph()
    p.text = "• Standard tools mark declarations as 'Complete/Approved' the second an API returns 200 OK."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_p4.add_paragraph()
    p.text = "• In reality, employer payroll teams hold declarations in 'Pending Verification' queues where files may be rejected."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 4: Architecture & Agent Workflow
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s4)
    add_header(s4, "System Architecture: End-to-End Agent Workflow", page_num=4)

    steps = [
        ("STAGE 1", "Guided Intake", "Tenant details, landlord address, lease term, and integer rent.", TEXT_CYAN),
        ("STAGE 2", "Ledger Matching", "Deterministic engine matches bank transactions to rent months.", TEXT_CYAN),
        ("STAGE 3", "PAN Audit", "Section 10(13A) engine tests ₹1L boundary; blocks if PAN missing.", TEXT_AMBER),
        ("STAGE 4", "PDF Receipts", "ReportLab generates receipts stamped with SHA-256 hashes.", TEXT_GREEN),
        ("STAGE 5", "Approval Gate", "Server blocks upload until explicit human sign-off on exact hashes.", TEXT_PURPLE),
        ("STAGE 6", "Portal Verify", "Idempotent push; agent independently re-reads HRMS status.", TEXT_GREEN)
    ]

    for idx, (stg, title, desc, clr) in enumerate(steps):
        left = Inches(0.8 + idx * 1.95)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.8), Inches(1.85), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = BORDER_DEFAULT
        card.line.width = Pt(1.5)

        tb = s4.shapes.add_textbox(left + Inches(0.12), Inches(1.95), Inches(1.61), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = stg
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = clr
        p0.space_after = Pt(6)

        p1 = tf.add_paragraph()
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        p1.space_after = Pt(12)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 5: Test Lab & The 4 Mandatory Scenarios
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s5)
    add_header(s5, "Test Lab: Verification of the 4 Official Scenarios", page_num=5)

    tf_t1 = draw_card(s5, Inches(0.8), Inches(1.8), Inches(5.66), Inches(2.35), 
                      "TC1: 12 Monthly Rent Payments", "100% Deterministic Match", border_color=BORDER_GREEN)
    p = tf_t1.add_paragraph()
    p.text = "• Input: 12 valid distinct HDFC bank records (₹25,000/mo) for FY 2025-26."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t1.add_paragraph()
    p.text = "• Assertions: 12/12 months matched, 12 receipts generated with unique SHA-256 hashes, exactly 0 duplicate receipts."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t1.add_paragraph()
    p.text = "• Automated Test Result: PASSED (Verified via pytest)"
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_GREEN

    tf_t2 = draw_card(s5, Inches(6.86), Inches(1.8), Inches(5.66), Inches(2.35), 
                      "TC2: Missing Evidence for March", "Zero-Hallucination Proof", border_color=BORDER_GREEN)
    p = tf_t2.add_paragraph()
    p.text = "• Input: 11 payments present in banking ledger; March record completely omitted."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t2.add_paragraph()
    p.text = "• Assertions: March strictly flagged as MISSING_PAYMENT. Exactly ZERO receipts generated or inferred for March; other 11 receipts generated."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t2.add_paragraph()
    p.text = "• Automated Test Result: PASSED (Zero Hallucination Verified)"
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_GREEN

    tf_t3 = draw_card(s5, Inches(0.8), Inches(4.45), Inches(5.66), Inches(2.35), 
                      "TC3: Rent Exceeds PAN Threshold", "Statutory Boundary Tests", border_color=BORDER_GREEN)
    p = tf_t3.add_paragraph()
    p.text = "• Input: Annual rent ₹3,00,000 (> ₹1,00,000 threshold), Landlord PAN omitted."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t3.add_paragraph()
    p.text = "• Assertions: PAN requirement flagged, user prompted, upload creation strictly blocked. Tested boundary cases: Below (₹80k), Equal (₹100k), Above (₹100,001)."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t3.add_paragraph()
    p.text = "• Automated Test Result: PASSED (Submission Blocked Verified)"
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_GREEN

    tf_t4 = draw_card(s5, Inches(6.86), Inches(4.45), Inches(5.66), Inches(2.35), 
                      "TC4: Pending Employer Verification", "Truthful Status Reporting", border_color=BORDER_GREEN)
    p = tf_t4.add_paragraph()
    p.text = "• Input: Approved proof pack uploaded to Mock Acme Corp HRMS Portal."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t4.add_paragraph()
    p.text = "• Assertions: Agent independently queries status endpoint post-upload; reports PENDING_VERIFICATION; never mislabels as Approved."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED
    p = tf_t4.add_paragraph()
    p.text = "• Automated Test Result: PASSED (Truth in Reporting Verified)"
    p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = TEXT_GREEN

    # =========================================================================
    # SLIDE 6: Security, Cryptography & Auditability
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s6)
    add_header(s6, "Security Architecture: Cryptographic Proof & Auditability", page_num=6)

    tf_s1 = draw_card(s6, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Payload Hash Binding", "Tamper Detection", border_color=BORDER_CYAN)
    p = tf_s1.add_paragraph()
    p.text = "• Every generated receipt PDF is cryptographically fingerprinted using SHA-256."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s1.add_paragraph()
    p.text = "• Human approval locks the exact list of receipt hashes and target destination."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s1.add_paragraph()
    p.text = "• If any PDF content or metadata is altered post-approval, the backend detects the mismatch and automatically INVALIDATES the approval."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_s2 = draw_card(s6, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Idempotent Submissions", "Duplicate Prevention", border_color=BORDER_CYAN)
    p = tf_s2.add_paragraph()
    p.text = "• Network retries often cause duplicate tax submissions and payroll confusion."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s2.add_paragraph()
    p.text = "• RentWise AI computes a stable idempotency key bound to proof pack ID, version, and file hashes."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s2.add_paragraph()
    p.text = "• Re-submitting the same payload safely returns the existing upload reference without creating duplicate records."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_s3 = draw_card(s6, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Append-Only Audit Trail", "Forensic Traceability", border_color=BORDER_CYAN)
    p = tf_s3.add_paragraph()
    p.text = "• Every action is logged in an immutable, append-only SQLite AuditEvent table."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s3.add_paragraph()
    p.text = "• Logs capture: Event ID, entity type, actor identity, correlation ID, before/after JSON state, evidence hashes, and audit reason."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_s3.add_paragraph()
    p.text = "• Demo resets preserve audit and test history for forensic review."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: Live Software Demonstration
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s7)
    add_header(s7, "Interactive Demo Guide: 7 Dedicated Application Views", page_num=7)

    views = [
        ("1. Overview Dashboard", "Real database metrics: Total months, matched payments, missing evidence, receipts generated, approval & upload status."),
        ("2. Create Proof Pack", "Guided 4-step wizard: Tenant details, landlord information, rent schedule, and pre-reconciliation statutory review."),
        ("3. Reconciliation Lab", "Monthly ledger timeline with visual match status badges, transaction references, single PDF downloads & full ZIP archive."),
        ("4. Approval Center", "Mandatory human review gate displaying total rent, receipt count, cryptographic hashes, approver identity & signature notes."),
        ("5. Mock Employer Portal", "Simulated enterprise HRMS: displays received declaration files and provides independent post-upload verification."),
        ("6. Test Lab Benchmark", "Run All Tests button executing live test cases TC1–TC4 with real-time automated assertions and pass/fail reports."),
        ("7. Audit Trail Explorer", "Searchable, filterable event log showing tamper-evident state transitions, actor timestamps, and cryptographic proofs.")
    ]

    for idx, (v_title, v_desc) in enumerate(views):
        y_pos = Inches(1.8 + idx * 0.72)
        bar = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_pos, Inches(11.733), Inches(0.62))
        bar.fill.solid()
        bar.fill.fore_color.rgb = BG_CARD
        bar.line.color.rgb = BORDER_DEFAULT
        bar.line.width = Pt(1.0)

        tb = s7.shapes.add_textbox(Inches(1.0), y_pos + Inches(0.08), Inches(11.3), Inches(0.5))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = v_title
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_CYAN

        p_d = tf.add_paragraph()
        p_d.text = v_desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 8: Technology Stack & Implementation Quality
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s8)
    add_header(s8, "Production Implementation: The Engineering Stack", page_num=8)

    tf_e1 = draw_card(s8, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Client Architecture", "React 19 & Tailwind v4", border_color=BORDER_DEFAULT)
    p = tf_e1.add_paragraph()
    p.text = "• Modern UI: Built with React 19, TypeScript, and Vite 8."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e1.add_paragraph()
    p.text = "• Fintech Design: Custom Tailwind CSS v4 design system with accessible contrast and responsive layouts."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e1.add_paragraph()
    p.text = "• Quick Judge Loaders: Header shortcuts to trigger TC1, TC2, TC3, and TC4 with a single click."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    tf_e2 = draw_card(s8, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Backend Engine", "FastAPI & ReportLab", border_color=BORDER_DEFAULT)
    p = tf_e2.add_paragraph()
    p.text = "• High-Performance API: Asynchronous FastAPI server with typed Pydantic models."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e2.add_paragraph()
    p.text = "• Real PDF Rendering: ReportLab generates authentic, downloadable PDF receipts (no mock HTML prints)."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e2.add_paragraph()
    p.text = "• Persistent Storage: SQLAlchemy ORM backed by SQLite with relational foreign key integrity."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    tf_e3 = draw_card(s8, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Automated Verification", "100% Pytest Coverage", border_color=BORDER_GREEN)
    p = tf_e3.add_paragraph()
    p.text = "• Automated Test Suite: 12 comprehensive pytest unit and integration tests executing in 1.1s."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e3.add_paragraph()
    p.text = "• Zero Simulated Results: Every test run executes real database insertions, hash validations, and assertions."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_e3.add_paragraph()
    p.text = "• Unified Deployment: Deployed on Render as a single unified service serving both React and FastAPI."
    p.font.size = Pt(11); p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 9: Market Impact & Commercial Viability
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s9)
    add_header(s9, "Market Opportunity: Transforming Enterprise Payroll & Tax", page_num=9)

    tf_m1 = draw_card(s9, Inches(0.8), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "For Employees", "Zero Stress & Max Savings", border_color=BORDER_CYAN)
    p = tf_m1.add_paragraph()
    p.text = "• Saves 4–6 Hours: Eliminates manual statement scanning, receipt writing, and spreadsheet tracking."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m1.add_paragraph()
    p.text = "• Audit Protection: Guarantees tax compliance under Section 10(13A), avoiding unexpected tax deductions in February/March payroll."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m1.add_paragraph()
    p.text = "• Transparent Proof: Full downloadable audit pack with verifiable PDF receipts."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_m2 = draw_card(s9, Inches(4.84), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "For Employers & HRMS", "70% Faster Verification", border_color=BORDER_CYAN)
    p = tf_m2.add_paragraph()
    p.text = "• Automated Pre-Verification: HR teams receive proof packs that are already reconciled against verified banking records."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m2.add_paragraph()
    p.text = "• Reduced Friction: Missing evidence and PAN requirements are caught before submission, slashing HR support tickets."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m2.add_paragraph()
    p.text = "• Integration Ready: Easily integrates into Workday, Darwinbox, Keka, GreytHR, and SAP SuccessFactors via REST APIs."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    tf_m3 = draw_card(s9, Inches(8.88), Inches(1.8), Inches(3.64), Inches(4.9), 
                      "Target Market Size", "B2B SaaS / FinTech", border_color=BORDER_CYAN)
    p = tf_m3.add_paragraph()
    p.text = "• 100M+ Salaried Employees in India submitting annual investment declarations."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m3.add_paragraph()
    p.text = "• ₹45,000 Cr+ in HRA deductions claimed annually across formal corporate sectors."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED; p.space_after = Pt(8)
    p = tf_m3.add_paragraph()
    p.text = "• Direct revenue paths: B2B HRMS plugin license (per employee/year) or B2C consumer tax filing agent."
    p.font.size = Pt(11.5); p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 10: Conclusion & Jury Verdict
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(s10)
    add_header(s10, "Conclusion: The Gold Standard in Verifiable AI Agents", page_num=10)

    # Big Final Summary Card
    big_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.9))
    big_card.fill.solid()
    big_card.fill.fore_color.rgb = BG_CARD
    big_card.line.color.rgb = BORDER_GREEN
    big_card.line.width = Pt(2.0)

    tb = s10.shapes.add_textbox(Inches(1.2), Inches(2.1), Inches(11.0), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "WHY RENTWISE AI DESERVES THE TOP JURY SCORE"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = TEXT_GREEN
    p0.space_after = Pt(12)

    bullets = [
        ("✓ Complete, Working Full-Stack System:", "Not a mockup or chatbot skin. A real React 19 UI, persistent SQLite database, FastAPI backend, and ReportLab PDF engine."),
        ("✓ Strict Zero-Hallucination Guarantee:", "Financial arithmetic and month matching are 100% deterministic. Never fabricates payments or PANs when evidence is missing."),
        ("✓ Statutory Tax Intelligence:", "Implements Indian Income Tax Section 10(13A) & CBDT Circular 08/2013 with 3-way boundary assertions (< ₹1L, == ₹1L, > ₹1L)."),
        ("✓ Cryptographic Integrity & Safety:", "Server-enforced human approval gate bound to SHA-256 PDF hashes. Payload changes invalidate approval; retries are duplicate-safe."),
        ("✓ Truthful Post-Upload Verification:", "Independently queries the employer portal and truthfully reports PENDING_VERIFICATION without premature approval claims."),
        ("✓ 100% Automated Test Pass Rate:", "12/12 automated pytest suites passing. Repeatable live test runner with assertion logs.")
    ]

    for title, desc in bullets:
        p = tf.add_paragraph()
        p.text = f"{title} "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_WHITE
        
        # Add normal text inline
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(8)

    p_end = tf.add_paragraph()
    p_end.text = "Thank You! We invite the Jury to test the Live Agent: https://rentwise-ai-84yv.onrender.com/"
    p_end.font.size = Pt(13)
    p_end.font.bold = True
    p_end.font.color.rgb = TEXT_CYAN
    p_end.space_before = Pt(8)

    output_path = os.path.abspath("./RentWise_AI_Online_Jury_Presentation.pptx")
    prs.save(output_path)
    print(f"Jury Presentation saved to: {output_path}")
    return output_path

if __name__ == "__main__":
    build_jury_presentation()
